import os
import re
from pathlib import Path
from sqlalchemy.orm import Session
from backend.app.models import Document, DocumentChunk, Course, User
from backend.app.config import settings

_SENTENCE_SPLIT = re.compile(r"(?<=[.!?…])\s+")

def chunk_text(text: str, chunk_size: int = 600, overlap: int = 80) -> list[str]:
    """Splits text into chunks with overlap, respecting sentence and paragraph boundaries (adapted from PaperBrain)."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    sentences = []
    for para in paragraphs:
        sentences.extend([s.strip() for s in _SENTENCE_SPLIT.split(para) if s.strip()])

    if not sentences:
        return []

    chunks = []
    current = ""
    for s in sentences:
        candidate = f"{current} {s}".strip() if current else s
        if len(candidate) <= chunk_size:
            current = candidate
            continue

        if current:
            chunks.append(current.strip())
            current = (current[-overlap:] + " " + s).strip() if overlap else s
        else:
            current = s

        while len(current) > chunk_size * 2:
            head, current = current[:chunk_size], current[chunk_size - overlap:]
            chunks.append(head.strip())

    if current:
        chunks.append(current.strip())

    return [c for c in chunks if c]

class DocumentService:
    @classmethod
    def extract_text_from_file(cls, file_path: str, ext: str) -> tuple[str, int]:
        """Extracts text and page count from PDF, DOCX, TXT, MD."""
        ext = ext.lower().lstrip(".")
        
        if ext == "pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                pages_text = []
                for idx, page in enumerate(reader.pages):
                    t = page.extract_text() or ""
                    if t.strip():
                        pages_text.append(f"--- [Page {idx + 1}] ---\n{t.strip()}")
                return "\n\n".join(pages_text), len(reader.pages)
            except Exception as e:
                return f"Error reading PDF: {e}", 1

        elif ext in ("docx", "doc"):
            try:
                import docx
                doc = docx.Document(file_path)
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                return "\n\n".join(paras), max(1, len(paras) // 10)
            except Exception as e:
                return f"Error reading DOCX: {e}", 1

        elif ext in ("txt", "md", "csv"):
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                return content, 1
            except Exception as e:
                return f"Error reading text file: {e}", 1

        elif ext in ("pptx", "ppt"):
            # Simple text extraction for PPTX
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(file_path, 'r') as z:
                    slide_texts = []
                    slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")])
                    for idx, s in enumerate(slide_files):
                        xml_content = z.read(s)
                        tree = ET.fromstring(xml_content)
                        texts = [node.text for node in tree.iter() if node.text]
                        slide_texts.append(f"--- [Slide {idx+1}] ---\n" + " ".join(texts))
                return "\n\n".join(slide_texts), max(1, len(slide_files))
            except Exception:
                return "Presentation content", 1

        return "", 1

    @classmethod
    def ingest_document(cls, db: Session, user: User, course_id: int, filename: str, file_bytes: bytes) -> Document:
        ext = Path(filename).suffix.lower().lstrip(".")
        course = db.query(Course).filter_by(id=course_id, user_id=user.id).first()
        if not course:
            raise ValueError("Course not found")

        # Save to disk
        course_dir = settings.UPLOAD_DIR / f"course_{course_id}"
        course_dir.mkdir(parents=True, exist_ok=True)
        file_path = course_dir / filename

        with open(file_path, "wb") as f:
            f.write(file_bytes)

        full_text, page_count = cls.extract_text_from_file(str(file_path), ext)

        doc = Document(
            user_id=user.id,
            course_id=course.id,
            filename=filename,
            file_type=ext,
            file_path=str(file_path),
            file_size=len(file_bytes),
            page_count=page_count
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        # Chunk and index
        chunks = chunk_text(full_text)
        for idx, chunk_content in enumerate(chunks):
            # Extract page number if present in chunk
            page_match = re.search(r"\[(?:Page|Slide)\s+(\d+)\]", chunk_content)
            page_num = int(page_match.group(1)) if page_match else 1

            chunk_record = DocumentChunk(
                document_id=doc.id,
                course_id=course.id,
                chunk_index=idx,
                content=chunk_content,
                page_number=page_num,
                token_count=len(chunk_content.split())
            )
            db.add(chunk_record)

        db.commit()
        return doc

